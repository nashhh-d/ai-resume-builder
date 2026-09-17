import io
import base64
import subprocess

from docx import Document
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from jinja2 import Environment


def build_docx(data, summary, suggested_skills):
    doc = Document()
    doc.add_heading(data.get("name", ""), 0)
    contact = f'{data.get("email","")} | {data.get("phone","")} | {data.get("linkedin","")}'
    doc.add_paragraph(contact)
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(summary)
    all_skills = data.get("skills", []) + suggested_skills
    if all_skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(all_skills))
    experience = data.get("experience", [])
    if experience:
        doc.add_heading("Experience", level=1)
        for exp in experience:
            role_line = f'{exp.get("role","")} - {exp.get("company","")} ({exp.get("duration","")})'
            doc.add_paragraph(role_line, style="List Bullet")
            description = exp.get("description", "")
            if description:
                doc.add_paragraph(description)
    education = data.get("education", [])
    if education:
        doc.add_heading("Education", level=1)
        for edu in education:
            line = f'{edu.get("degree","")} - {edu.get("institution","")} ({edu.get("year","")})'
            doc.add_paragraph(line, style="List Bullet")
    projects = data.get("projects", [])
    if projects:
        doc.add_heading("Projects", level=1)
        for proj in projects:
            line = f'{proj.get("title","")}: {proj.get("description","")}'
            doc.add_paragraph(line, style="List Bullet")
    certs = data.get("certifications", [])
    if certs:
        doc.add_heading("Certifications", level=1)
        for cert in certs:
            doc.add_paragraph(cert, style="List Bullet")
    buffer = io.BytesIO()
    doc.save(buffer)
    return base64.b64encode(buffer.getvalue()).decode()


def build_pdf(data, summary, suggested_skills):
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer)
    styles = getSampleStyleSheet()
    elements = []
    elements.append(Paragraph(data.get("name",""), styles["Title"]))
    contact = f'{data.get("email","")} | {data.get("phone","")} | {data.get("linkedin","")}'
    elements.append(Paragraph(contact, styles["Normal"]))
    elements.append(Spacer(1, 10))
    elements.append(Paragraph("<b>Summary</b>", styles["Heading2"]))
    elements.append(Paragraph(summary, styles["Normal"]))
    elements.append(Spacer(1, 10))
    all_skills = data.get("skills", []) + suggested_skills
    if all_skills:
        elements.append(Paragraph("<b>Skills</b>", styles["Heading2"]))
        elements.append(Paragraph(", ".join(all_skills), styles["Normal"]))
        elements.append(Spacer(1, 10))
    experience = data.get("experience", [])
    if experience:
        elements.append(Paragraph("<b>Experience</b>", styles["Heading2"]))
        for exp in experience:
            role_line = f'{exp.get("role","")} - {exp.get("company","")} ({exp.get("duration","")})'
            elements.append(Paragraph(role_line, styles["Normal"]))
            description = exp.get("description", "")
            if description:
                elements.append(Paragraph(description, styles["Normal"]))
        elements.append(Spacer(1, 10))
    education = data.get("education", [])
    if education:
        elements.append(Paragraph("<b>Education</b>", styles["Heading2"]))
        for edu in education:
            line = f'{edu.get("degree","")} - {edu.get("institution","")} ({edu.get("year","")})'
            elements.append(Paragraph(line, styles["Normal"]))
        elements.append(Spacer(1, 10))
    projects = data.get("projects", [])
    if projects:
        elements.append(Paragraph("<b>Projects</b>", styles["Heading2"]))
        for proj in projects:
            line = f'{proj.get("title","")}: {proj.get("description","")}'
            elements.append(Paragraph(line, styles["Normal"]))
        elements.append(Spacer(1, 10))
    certs = data.get("certifications", [])
    if certs:
        elements.append(Paragraph("<b>Certifications</b>", styles["Heading2"]))
        for cert in certs:
            elements.append(Paragraph(cert, styles["Normal"]))
    doc.build(elements)
    return base64.b64encode(buffer.getvalue()).decode()


def build_latex_resume(data):
    name = data.get("name", "")
    email = data.get("email", "")
    phone = data.get("phone", "")
    linkedin = data.get("linkedin", "")
    github = data.get("github", "")
    education = data.get("education", [])
    experience = data.get("experience", [])
    projects = data.get("projects", [])
    skills = data.get("skills", [])
    certifications = data.get("certifications", [])

    # Use << >> delimiters to avoid conflict with LaTeX { }
    env = Environment(
        variable_start_string='<<',
        variable_end_string='>>',
        block_start_string='<%',
        block_end_string='%>',
    )

    latex_template = r"""
\documentclass[letterpaper,11pt]{article}

\usepackage[empty]{fullpage}
\usepackage{titlesec}
\usepackage{enumitem}
\usepackage[hidelinks]{hyperref}
\usepackage{fancyhdr}
\usepackage{tabularx}

\pagestyle{fancy}
\fancyhf{}
\renewcommand{\headrulewidth}{0pt}

\begin{document}

\begin{center}
    {\Huge \scshape << name >> } \\ \vspace{1pt}
    \small
    << phone >> $\cdot$ \href{mailto:<< email >>}{<< email >>} $\cdot$ \href{https://linkedin.com/in/<< linkedin >>}{linkedin.com/in/<< linkedin >>} $\cdot$ \href{https://github.com/<< github >>}{github.com/<< github >>}
\end{center}

\section*{Education}

<% for edu in education %>
\textbf{<< edu.degree >>} \\
<< edu.institution >> (<< edu.year >>)
\vspace{4pt}

<% endfor %>

\section*{Experience}

<% for exp in experience %>
\textbf{<< exp.role >>} — << exp.company >> (<< exp.duration >>)
\begin{itemize}[leftmargin=*]
\item << exp.description >>
\end{itemize}

<% endfor %>

\section*{Projects}

<% for proj in projects %>
\textbf{<< proj.title >>}
\begin{itemize}[leftmargin=*]
\item << proj.description >>
\end{itemize}

<% endfor %>

\section*{Skills}

<< skills | join(", ") >>

<% if certifications %>
\section*{Certifications}

<% for cert in certifications %>
- << cert >> \\
<% endfor %>

<% endif %>

\end{document}
"""

    template = env.from_string(latex_template)
    rendered_latex = template.render(
        name=name, email=email, phone=phone,
        linkedin=linkedin, github=github,
        education=education, experience=experience,
        projects=projects, skills=skills,
        certifications=certifications
    )

    import tempfile, os
    with tempfile.TemporaryDirectory() as tmpdir:
        tex_file = os.path.join(tmpdir, "resume.tex")
        pdf_file = os.path.join(tmpdir, "resume.pdf")

        with open(tex_file, "w", encoding="utf-8") as f:
            f.write(rendered_latex)

        result = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "-output-directory", tmpdir, tex_file],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE
        )

        if not os.path.exists(pdf_file):
            raise RuntimeError(
                f"pdflatex failed to produce a PDF.\n"
                f"STDOUT: {result.stdout.decode(errors='ignore')}\n"
                f"STDERR: {result.stderr.decode(errors='ignore')}"
            )

        with open(pdf_file, "rb") as f:
            pdf_bytes = f.read()

    return base64.b64encode(pdf_bytes).decode()
